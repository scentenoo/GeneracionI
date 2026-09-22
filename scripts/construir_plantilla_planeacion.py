"""Genera templates/planeacion_individual.docx con el formato Diario
Pedagógico del programa (PC-PA-003-F03), calcado del formato oficial en
blanco que pasó la Secretaría (PROPUESTA_FORMATO.pdf): tabla de MOMENTOS /
TIEMPO / ACTIVIDAD con una fila por momento (inicial, desarrollo, cierre) y
un único campo de evaluación de la clase — ya no hay columnas separadas de
reflexión pedagógica y avances/retrocesos, esas se eliminaron porque el
formato oficial no las tiene. Ver el contrato de variables en
client/src/ui/planeacion_screen._contexto_documento.

    python scripts/construir_plantilla_planeacion.py

Cuidado con los {%tr ...%} de la tabla de asistencia: el for y el endfor
van en filas SEPARADAS, o docxtpl tira "unknown tag 'endfor'".

El membrete (encabezado y pie de página) sale de _membrete_municipio.py,
compartido con construir_plantilla_gestion.py. Acá copia el formato oficial
real —CÓDIGO PC-PA-003-F03, VERSIÓN 0—: cada sección (datos de la clase,
momentos, evaluación, asistencia, fotos) va en su propio recuadro con
borde, igual que el formato oficial. También hizo falta fijar la fuente a
mano: sin eso, el texto sale en la fuente "menor" del tema de Word
(Cambria, con serifa) en vez de la fuente sin serifa del resto de
documentos del programa — de ahí _fijar_fuente al final.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "templates" / "planeacion_individual.docx"

FUENTE = "Calibri"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _membrete_municipio import armar_membrete_  # noqa: E402


def celda_titulo(celda, texto):
    p = celda.paragraphs[0]
    r = p.add_run(texto)
    r.bold = True


def _ancho_columnas(tabla, anchos):
    """python-docx necesita el ancho puesto en la columna Y en cada celda
    de cada fila — con solo `tabla.columns[i].width` Word igual autoajusta
    y las columnas quedan destimbaladas entre esta tabla y la de al lado."""
    tabla.autofit = False
    for col, ancho in zip(tabla.columns, anchos):
        col.width = ancho
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = ancho


_ANCHO_PAGINA = Cm(18)


def _recuadro(doc, texto=None):
    """Recuadro con borde propio (una tabla de 1x1) del ancho de toda la
    página — el mismo recurso visual que el documento real usa para
    "ASISTENCIA" y "EVIDENCIA FOTOGRÁFICA DE LA CLASE", en vez de un párrafo
    suelto. Sin el ancho explícito la tabla nace angosta (autofit por
    contenido) y las fotos de la evidencia, en vez de entrar una al lado
    de la otra, se apilan y se salen de la página."""
    caja = doc.add_table(rows=1, cols=1)
    caja.style = "Table Grid"
    _ancho_columnas(caja, [_ANCHO_PAGINA])
    if texto is not None:
        p = caja.rows[0].cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold = True
        r.italic = True
    return caja


def _fijar_fuente(doc, nombre=FUENTE):
    """Sin esto, cualquier texto que no pase por un run con fuente propia
    hereda la fuente "menor" del tema de Word (Cambria, con serifa) en vez
    de una fuente sin serifa — se nota fuerte comparado con el resto de
    documentos del programa. Recorre cuerpo, encabezado y pie, tablas
    incluidas (y tablas anidadas, por si alguna celda llega a tener una).
    No toca un run que ya tenga su propia fuente puesta a mano (el
    membrete institucional y el pie de página van siempre en Century
    Gothic — ver _FUENTE_INSTITUCION en _membrete_municipio.py)."""

    def en_parrafos(parrafos):
        for p in parrafos:
            for r in p.runs:
                if r.font.name is None:
                    r.font.name = nombre

    def en_tablas(tablas):
        for tabla in tablas:
            for fila in tabla.rows:
                for celda in fila.cells:
                    en_parrafos(celda.paragraphs)
                    en_tablas(celda.tables)

    contenedores = [doc] + [s.header for s in doc.sections] + [s.footer for s in doc.sections]
    for contenedor in contenedores:
        en_parrafos(contenedor.paragraphs)
        en_tablas(contenedor.tables)

    doc.styles["Normal"].font.name = nombre


def main():
    doc = Document()
    armar_membrete_(doc, "DIARIO PEDAGÓGICO", codigo_tramite="PC-PA-003-F03", version_tramite="0")

    # Datos de la clase: FECHA / GRUPO / OBJETIVO / TEMAS DE LA CLASE.
    tabla = doc.add_table(rows=4, cols=3)
    tabla.style = "Table Grid"
    _ancho_columnas(tabla, [Cm(6), Cm(6), Cm(6)])

    filas_kv = [
        ("FECHA", "{{ fecha }}"),
        ("GRUPO", "{{ grupo }}"),
        ("OBJETIVO", "{{ objetivo }}"),
        ("TEMAS DE LA CLASE",
         "{% for t in temas_vistos %}{{ t }}{% if not loop.last %}, {% endif %}{% endfor %}"),
    ]
    for i, (k, v) in enumerate(filas_kv):
        celda_titulo(tabla.rows[i].cells[0], k)
        valor = tabla.rows[i].cells[1].merge(tabla.rows[i].cells[2])
        valor.paragraphs[0].add_run(v)

    # Momentos de la clase: MOMENTOS / TIEMPO / ACTIVIDAD, una fila por
    # momento — igual que la tabla del formato oficial en blanco.
    doc.add_paragraph()
    momentos_tabla = doc.add_table(rows=4, cols=3)
    momentos_tabla.style = "Table Grid"
    _ancho_columnas(momentos_tabla, [Cm(3.5), Cm(2.5), Cm(12)])

    fila_cabecera = momentos_tabla.rows[0]
    for celda, texto in zip(fila_cabecera.cells, ["MOMENTOS", "TIEMPO", "ACTIVIDAD"]):
        celda.paragraphs[0].add_run(texto).bold = True

    momentos = [
        ("MOMENTO INICIAL", "{{ momento_inicial_min }}", "{{ momento_inicial_texto }}"),
        ("MOMENTO DE DESARROLLO", "{{ momento_desarrollo_min }}", "{{ momento_desarrollo_texto }}"),
        ("MOMENTO DE CIERRE", "{{ momento_final_min }}", "{{ momento_final_texto }}"),
    ]
    for fila, (etiqueta, min_ph, texto_ph) in zip(momentos_tabla.rows[1:], momentos):
        fila.cells[0].paragraphs[0].add_run(etiqueta).bold = True
        fila.cells[1].paragraphs[0].add_run(f"{min_ph} minutos")
        fila.cells[2].paragraphs[0].add_run(texto_ph)

    # Evaluación de la clase: un único campo — el formato oficial ya no
    # separa reflexión pedagógica y avances/retrocesos como la versión
    # anterior de esta plantilla. Solo el título del recuadro (que ya es el
    # rótulo del campo); no se repite "Observaciones del desempeño de los
    # estudiantes" como subtítulo, para no mezclar instrucción con el
    # contenido real ya diligenciado.
    doc.add_paragraph()
    caja_eval = doc.add_table(rows=1, cols=1)
    caja_eval.style = "Table Grid"
    _ancho_columnas(caja_eval, [_ANCHO_PAGINA])
    celda_eval = caja_eval.rows[0].cells[0]
    celda_eval.paragraphs[0].add_run("EVALUACIÓN DE LA CLASE").bold = True
    celda_eval.add_paragraph("{{ observaciones }}")

    # Evidencias: un recuadro por cada evidencia (asistencia, fotos) —
    # igual que el documento real, en vez de títulos sueltos sin borde. Los
    # títulos de cada recuadro son los rótulos exactos del formato oficial
    # (ASISTENCIA / EVIDENCIA FOTOGRÁFICA DE LA CLASE), sin un encabezado
    # "EVIDENCIAS" extra que el formato oficial no tiene.
    doc.add_paragraph()

    _recuadro(doc, "ASISTENCIA")
    asis = doc.add_table(rows=1, cols=2)
    asis.style = "Table Grid"
    _ancho_columnas(asis, [Cm(9), Cm(9)])
    asis.rows[0].cells[0].paragraphs[0].add_run("ESTUDIANTE").bold = True
    asis.rows[0].cells[1].paragraphs[0].add_run("Asistió").bold = True
    fila_for = asis.add_row().cells
    fila_for[0].paragraphs[0].add_run("{%tr for a in asistencia %}")
    fila_datos_asis = asis.add_row().cells
    fila_datos_asis[0].paragraphs[0].add_run("{{ a.nombre }}")
    fila_datos_asis[1].paragraphs[0].add_run("{{ a.presente }}")
    fila_endfor = asis.add_row().cells
    fila_endfor[0].paragraphs[0].add_run("{%tr endfor %}")

    # Fotos (1 a 3): un solo run con el for/endfor adentro, para que
    # docxtpl no lo vea partido entre runs distintos (eso rompería el tag,
    # igual que con la tabla de asistencia de más arriba).
    doc.add_paragraph()
    _recuadro(doc, "EVIDENCIA FOTOGRÁFICA DE LA CLASE")
    caja_fotos = _recuadro(doc)
    pf = caja_fotos.rows[0].cells[0].paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.add_run("{% for foto in fotos_clase %}{{ foto }} {% endfor %}")

    _fijar_fuente(doc)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(SALIDA))
    print(f"Plantilla generada en {SALIDA}")


if __name__ == "__main__":
    main()
