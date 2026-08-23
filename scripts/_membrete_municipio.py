"""Encabezado y pie de página oficiales del municipio, compartidos por los
scripts que arman las plantillas (construir_plantilla_planeacion.py,
construir_plantilla_gestion.py).

El pie de página (web, dirección, correo) es el mismo en cualquier trámite
del municipio — no depende de cuál formato sea. El encabezado sí cambia por
trámite: el título y, cuando existe, el CÓDIGO/VERSIÓN de ese formato
puntual. Para "horas de gestión" no hay un documento oficial real del que
copiar un código (no existe, según confirmó el equipo directivo), así que
esa columna se omite en vez de inventar un número.

El escudo sale de un documento real (Diario Pedagógico, PC-PA-003-F03) que
pasó el equipo directivo — es el escudo público del municipio, no cambia
según el trámite, así que se reusa para cualquier formato.
"""

from __future__ import annotations

from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parents[1]
ESCUDO = RAIZ / "templates" / "assets" / "escudo_municipio.jpg"

_WEB_Y_DIRECCION = (
    "www.sanpedrodelosmilagros-antioquia.gov.co / Carrera 49A No.49-36 "
    "Parque principal - PBX. 8687039"
)
_CORREO_Y_CODIGO_POSTAL = (
    "E-mail: secgobierno@sanpedrodelosmilagros-antioquia.gov.co "
    "/ Código Postal 051010 / Página "
)


def agregar_campo_word_(paragraph, codigo_campo):
    """PAGE / NUMPAGES no son texto fijo: hay que insertar el campo de Word
    a mano, python-docx no tiene una función para esto."""
    run = paragraph.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = codigo_campo
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run._r.append(inicio)
    run._r.append(instruccion)
    run._r.append(fin)


def armar_membrete_(doc, titulo_documento, codigo_tramite=None, version_tramite=None):
    """Encabezado (escudo, título del documento, código/versión si el
    trámite tiene uno) y pie de página (web, dirección, correo, número de
    página) — igual al formato oficial que usa el programa.

    Sin `codigo_tramite`, la columna de código/versión no se dibuja: mejor
    dejarla afuera que inventar un número que no existe.
    """
    seccion = doc.sections[0]

    tiene_codigo = codigo_tramite is not None
    columnas = 4 if tiene_codigo else 3
    ancho_institucion = Cm(6.3) if tiene_codigo else Cm(7.5)
    ancho_titulo = Cm(6.4) if tiene_codigo else Cm(7.5)
    anchos = [ancho_institucion, ancho_titulo, Cm(2.5)]
    if tiene_codigo:
        anchos.append(Cm(2.7))

    encabezado = seccion.header
    encabezado.is_linked_to_previous = False
    tabla = encabezado.add_table(rows=3, cols=columnas, width=Cm(18))
    tabla.autofit = False
    for col, ancho in zip(tabla.columns, anchos):
        col.width = ancho
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = ancho

    # Institución y título: una sola celda por columna, fundida en las 3 filas.
    celda_institucion = tabla.cell(0, 0).merge(tabla.cell(2, 0))
    celda_institucion.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = celda_institucion.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run("ADMINISTRACIÓN MUNICIPAL")
    p2 = celda_institucion.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Secretaria de Educación y Desarrollo Social")
    r2.font.size = Pt(9)

    celda_titulo = tabla.cell(0, 1).merge(tabla.cell(2, 1))
    celda_titulo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    pt = celda_titulo.paragraphs[0]
    pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = pt.add_run(titulo_documento)
    rt.bold = True
    rt.font.size = Pt(12)

    celda_escudo = tabla.cell(0, 2).merge(tabla.cell(2, 2))
    celda_escudo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    pe = celda_escudo.paragraphs[0]
    pe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ESCUDO.exists():
        pe.add_run().add_picture(str(ESCUDO), width=Cm(2.1))

    if tiene_codigo:
        # Código y versión: van en filas distintas de la MISMA columna, no fundidas.
        p_codigo = tabla.cell(0, 3).paragraphs[0]
        p_codigo.add_run("CÓDIGO:").bold = True
        p_codigo_v = tabla.cell(0, 3).add_paragraph()
        p_codigo_v.add_run(codigo_tramite).bold = True
        if version_tramite is not None:
            p_version = tabla.cell(2, 3).paragraphs[0]
            p_version.add_run(f"VERSIÓN:  {version_tramite}").bold = True

    for fila in tabla.rows:
        for celda in fila.cells:
            for p in celda.paragraphs:
                for r in p.runs:
                    if r.font.size is None:
                        r.font.size = Pt(9)

    pie = seccion.footer
    pie.is_linked_to_previous = False
    p_web = pie.paragraphs[0]
    p_web.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_web = p_web.add_run(_WEB_Y_DIRECCION)
    r_web.font.size = Pt(8)

    p_mail = pie.add_paragraph()
    p_mail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mail.add_run(_CORREO_Y_CODIGO_POSTAL)
    agregar_campo_word_(p_mail, "PAGE")
    p_mail.add_run(" de ")
    agregar_campo_word_(p_mail, "NUMPAGES")
    for r in p_mail.runs:
        r.font.size = Pt(8)
