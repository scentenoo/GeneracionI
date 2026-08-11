"""Simulacro de un mes real con las 18 personas.

Ejerce el flujo completo contra el backend real: cada docente carga sus
planeaciones del mes con foto y asistencia, algunos registran otras
actividades, los directivos cargan horas de gestión, todos entregan su
informe, y el equipo directivo consolida (dashboard, descargas, ZIP).
También prueba a propósito las rutas de error y los límites de permisos.

Todo lo que sale distinto de lo esperado se anota como HALLAZGO. Al final
se limpia lo creado, para no dejar el mes de prueba en la Sheet.

    python scripts/simulacro_mes.py

Corre sobre un mes de prueba (MES_SIM) para no chocar con datos reales.
"""

from __future__ import annotations

import os
import re
import sys
import time
import tempfile
import traceback
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "client" / "src"))

from PIL import Image  # noqa: E402
import api_client  # noqa: E402
from services import image_utils, docx_generator  # noqa: E402

MES_SIM = "2026-09"          # mes abierto (cierra el 05/10, aún futuro)
MES_CERRADO = "2026-06"      # mes ya cerrado, para probar el bloqueo
ADMIN_USER, ADMIN_PASS = "samir", "GeneracionI2026!"
CREDS = RAIZ / "credenciales_iniciales.txt"

HALLAZGOS: list[str] = []
tiempos: list[tuple[str, float]] = []


def hallazgo(sev: str, texto: str):
    HALLAZGOS.append(f"[{sev}] {texto}")
    print(f"   ⚠ {sev}: {texto}")


def cron(nombre, fn, validar=None):
    """Corre fn con reintentos ante propagación (404/timeout) y mide.

    `validar` es un predicado opcional sobre el resultado: si devuelve
    False, se trata como respuesta malformada y se reintenta. Apps Script
    a veces contesta 200 con un cuerpo con forma incorrecta (una lista de
    strings donde esperábamos objetos); sin esto, el simulacro se caía con
    un AttributeError en vez de reintentar.
    """
    t0 = time.time()
    ultimo = None
    for intento in range(8):
        try:
            r = fn()
            if validar is not None and not validar(r):
                hallazgo("ALTA", f"{nombre}: respuesta malformada del backend "
                         f"(reintento {intento + 1}) — tipo {type(r).__name__}")
                time.sleep(5)
                continue
            tiempos.append((nombre, time.time() - t0))
            return r
        except api_client.ApiError as e:
            ultimo = e
            if "404" in str(e) or "tardando" in str(e) or "inesperado" in str(e):
                time.sleep(5)
                continue
            raise
    raise ultimo or RuntimeError(f"{nombre}: no se pudo tras reintentos")


def _lista_de_dicts(r):
    return isinstance(r, list) and all(isinstance(x, dict) for x in r)


def parse_credenciales() -> dict[str, str]:
    """usuario -> clave, del archivo de credenciales (varios bloques)."""
    if not CREDS.exists():
        raise SystemExit(f"No está {CREDS}. Corré cargar_nomina primero.")
    texto = CREDS.read_text(encoding="utf-8")
    creds = {}
    for m in re.finditer(r"usuario:\s*(\S+)\s*\n\s*clave:\s*(\S+)", texto):
        creds[m.group(1)] = m.group(2)
    return creds


def esperar(cond, seg=15):
    t0 = time.time()
    while time.time() - t0 < seg:
        if cond():
            return True
        time.sleep(0.3)
    return False


def main():
    print("=" * 60)
    print(f"SIMULACRO DE MES — {MES_SIM}")
    print("=" * 60)

    creds = parse_credenciales()
    print(f"\nCredenciales encontradas: {len(creds)} usuarios")

    foto_path = os.path.join(tempfile.gettempdir(), "sim_foto.jpg")
    Image.new("RGB", (60, 40), (80, 110, 90)).save(foto_path)
    foto_payload = image_utils.foto_a_payload(foto_path)

    def palabras(n):
        return " ".join(f"palabra{i}" for i in range(n))

    # --- login admin y estado inicial -----------------------------------
    print("\n--- LOGIN de las 18 personas ---")
    admin = cron("login admin", lambda: api_client.login(ADMIN_USER, ADMIN_PASS))
    tok_admin = admin["token"]

    usuarios = cron("listar_usuarios", lambda: api_client.listar_usuarios(tok_admin), _lista_de_dicts)
    cursos = cron("listar_cursos", lambda: api_client.listar_todos_los_cursos(tok_admin), _lista_de_dicts)
    activos = [c for c in cursos if c.get("activo")]
    por_id = {u["id"]: u for u in usuarios}

    sesiones = {}
    for u in usuarios:
        user = u["usuario"]
        if user == ADMIN_USER:
            sesiones[u["id"]] = admin
            continue
        clave = creds.get(user)
        if not clave:
            hallazgo("MEDIA", f"Sin credencial para {user} ({u['nombre']}): no puede entrar")
            continue
        try:
            sesiones[u["id"]] = cron(f"login {user}", lambda c=clave, x=user: api_client.login(x, c))
        except api_client.ApiError as e:
            hallazgo("ALTA", f"{user} no puede entrar: {e}")
    print(f"   entraron {len(sesiones)} de {len(usuarios)}")

    # --- SECCIÓN A: carga del mes ---------------------------------------
    print(f"\n--- A. Cada docente carga sus 4 clases de {MES_SIM} ---")
    dias = ["07", "14", "21", "28"]
    creadas = 0
    for c in activos:
        ses = sesiones.get(c["docente_id"])
        if not ses:
            hallazgo("MEDIA", f"Curso «{c['nombre']}» sin docente con sesión (id {c['docente_id']})")
            continue
        tok = ses["token"]
        for d in dias:
            datos = {
                "fecha": f"{MES_SIM}-{d}",
                "curso_id": c["id"],
                "objetivo": "Objetivo de la clase. " + palabras(25),
                "temas_vistos": ["Tema A", "Tema B"],
                "bloques": [{"momento": "Inicio", "minutos": 120,
                             "observacion": "Obs. " + palabras(25),
                             "avance": "Avance. " + palabras(25)}],
                "asistencia": [{"nombre": "Estudiante 1", "presente": True},
                               {"nombre": "Estudiante 2", "presente": False}],
            }
            try:
                cron(f"planeacion {c['id']} {d}",
                     lambda dd=datos, tk=tok: api_client.guardar_planeacion(tk, dd, {"foto_clase": foto_payload}))
                creadas += 1
            except api_client.ApiError as e:
                hallazgo("ALTA", f"No se pudo cargar clase {c['nombre']} {MES_SIM}-{d}: {e}")
    print(f"   planeaciones creadas: {creadas} (esperadas {len(activos) * 4})")

    # otras actividades (una por un par de docentes)
    print("\n--- A2. Otras actividades del mes ---")
    for c in activos[:3]:
        ses = sesiones.get(c["docente_id"])
        if not ses:
            continue
        datos = {"curso_id": c["id"], "fecha": f"{MES_SIM}-15",
                 "descripcion": "Reunión de docentes. " + palabras(10),
                 "horas_sede": 2, "horas_externas": 0}
        try:
            cron("actividad", lambda dd=datos, tk=ses["token"]:
                 api_client.guardar_actividad(tk, dd, {"foto": foto_payload}))
        except api_client.ApiError as e:
            hallazgo("MEDIA", f"No se pudo cargar actividad en {c['nombre']}: {e}")

    # horas de gestión (directivos)
    print("\n--- A3. Horas de gestión de los directivos ---")
    directivos = [u for u in usuarios if u["rol"] in ("directivo", "ambos")]
    for u in directivos:
        ses = sesiones.get(u["id"])
        if not ses:
            continue
        datos = {"fecha": f"{MES_SIM}-10", "actividad": "Coordinación del núcleo. " + palabras(8),
                 "horas_sede": "6", "entregable": "Acta"}
        try:
            cron("horas_gestion", lambda dd=datos, tk=ses["token"]:
                 api_client.guardar_horas_gestion(tk, dd))
        except api_client.ApiError as e:
            hallazgo("ALTA", f"{u['nombre']} no pudo cargar horas de gestión: {e}")

    # --- SECCIÓN B: rutas de error y permisos ---------------------------
    print("\n--- B. Validaciones y permisos (se ESPERA que rechacen) ---")

    def espera_error(desc, fn, debe_contener=None):
        try:
            fn()
            hallazgo("ALTA", f"{desc}: NO fue rechazado (debería))")
        except api_client.ApiError as e:
            if debe_contener and debe_contener.lower() not in str(e).lower():
                hallazgo("BAJA", f"{desc}: rechazado pero con mensaje raro: {e}")
            else:
                print(f"   ✓ {desc}: rechazado correctamente")

    # login malo
    espera_error("Login con clave incorrecta",
                 lambda: api_client.login(ADMIN_USER, "claveMala123"))

    # planeación con pocas palabras
    algun_curso = next((c for c in activos if sesiones.get(c["docente_id"])), None)
    if algun_curso:
        tok = sesiones[algun_curso["docente_id"]]["token"]
        espera_error("Planeación con objetivo corto (<20 palabras)",
                     lambda: api_client.guardar_planeacion(tok, {
                         "fecha": f"{MES_SIM}-07", "curso_id": algun_curso["id"],
                         "objetivo": "corto", "temas_vistos": ["x"],
                         "bloques": [{"momento": "a", "minutos": 120, "observacion": "corto", "avance": "corto"}],
                         "asistencia": [{"nombre": "E", "presente": True}],
                     }, {"foto_clase": foto_payload}))

    # docente viendo planeaciones de otro
    docentes = [u for u in usuarios if u["rol"] == "docente" and sesiones.get(u["id"])]
    if len(docentes) >= 2:
        d0, d1 = docentes[0], docentes[1]
        espera_error("Docente viendo planeaciones de otro docente",
                     lambda: api_client.obtener_planeaciones(sesiones[d0["id"]]["token"], docente_id=d1["id"]),
                     debe_contener="permiso")

    # docente entregando informe con mes incompleto (curso ajeno sin cargar)
    # (usamos un mes futuro sin planeaciones)
    if algun_curso:
        tok = sesiones[algun_curso["docente_id"]]["token"]
        espera_error("Informe de un mes sin planeaciones",
                     lambda: api_client.guardar_informe_mensual(
                         tok, algun_curso["id"], "2026-12",
                         {k: palabras(25) for k in ["objetivo_cumplimiento", "logros_avances",
                          "dificultades", "estrategias", "situacion_positiva", "ctei_integracion"]}))

    # un docente intentando editar horas de gestión ajenas / cargar gestión sin rol
    if docentes:
        espera_error("Docente puro cargando horas de gestión",
                     lambda: api_client.guardar_horas_gestion(sesiones[docentes[0]["id"]]["token"],
                         {"fecha": f"{MES_SIM}-10", "actividad": palabras(8), "horas_sede": "3"}),
                     debe_contener="rol")

    # --- SECCIÓN C: entrega de informes ---------------------------------
    print("\n--- C. Cada quien entrega su informe del mes ---")
    narr = {k: "Respuesta. " + palabras(25) for k in
            ["objetivo_cumplimiento", "logros_avances", "dificultades",
             "estrategias", "situacion_positiva", "ctei_integracion"]}
    narr["avance_semanal"] = []
    gest = {k: "Gestión. " + palabras(25) for k in
            ["objetivos", "logros", "novedades", "estrategias", "pendientes"]}

    entregados = 0
    # informe docente por curso (el dueño lo entrega; los "ambos" marcan gestión en el 1er curso)
    marcado_gestion = set()
    for c in activos:
        ses = sesiones.get(c["docente_id"])
        if not ses:
            continue
        u = por_id[c["docente_id"]]
        incluir = False
        gestion_narr = None
        if u["rol"] in ("directivo", "ambos") and c["docente_id"] not in marcado_gestion:
            incluir = True
            gestion_narr = gest
            marcado_gestion.add(c["docente_id"])
        try:
            cron(f"informe {c['id']}", lambda cc=c, tk=ses["token"], g=gestion_narr, inc=incluir:
                 api_client.guardar_informe_mensual(tk, cc["id"], MES_SIM, narr, g, inc))
            entregados += 1
        except api_client.ApiError as e:
            hallazgo("ALTA", f"No se pudo entregar informe de «{c['nombre']}»: {e}")
    print(f"   informes docentes entregados: {entregados} de {len(activos)}")

    # informe de gestión de directivos sin curso
    con_curso = {c["docente_id"] for c in activos}
    directivos_puros = [u for u in directivos if u["id"] not in con_curso and sesiones.get(u["id"])]
    for u in directivos_puros:
        try:
            cron("informe_gestion", lambda tk=sesiones[u["id"]]["token"]:
                 api_client.guardar_informe_gestion(tk, MES_SIM, gest))
            print(f"   informe de gestión entregado: {u['nombre']}")
        except api_client.ApiError as e:
            hallazgo("ALTA", f"{u['nombre']} no pudo entregar su informe de gestión: {e}")

    # --- SECCIÓN D: vista directiva -------------------------------------
    print("\n--- D. Consolidación del equipo directivo ---")
    dash = cron("dashboard", lambda: api_client.obtener_dashboard_directivo(tok_admin, MES_SIM))
    al_dia = sum(1 for d in dash if d["faltantes"] == 0)
    con_informe = sum(1 for d in dash if d["informe_entregado"])
    print(f"   dashboard: {len(dash)} cursos, {al_dia} con clases completas, {con_informe} con informe")
    if al_dia != len(dash):
        hallazgo("MEDIA", f"Dashboard: {len(dash) - al_dia} cursos sin las 4 clases tras la carga")
    if con_informe != len(dash):
        hallazgo("MEDIA", f"Dashboard: {len(dash) - con_informe} cursos sin informe tras la entrega")

    dirs_sc = cron("directivos_sin_curso", lambda:
                   api_client.directivos_sin_curso_del_mes(tok_admin, MES_SIM))
    print(f"   directivos sin curso: {len(dirs_sc)}, entregaron "
          f"{sum(1 for d in dirs_sc if d['informe_entregado'])}")

    # generar y validar el .docx de una muestra (docente y gestión)
    print("\n--- D2. Generar y verificar documentos ---")
    from docx import Document

    def verificar_docx(ruta, etiqueta):
        d = Document(ruta)
        texto = "\n".join(p.text for p in d.paragraphs)
        for tb in d.tables:
            for row in tb.rows:
                texto += "\n" + " | ".join(cel.text for cel in row.cells)
        sin_resolver = re.findall(r"{{.*?}}|{%.*?%}", texto)
        literal_none = "None" in re.findall(r"\bNone\b", texto)
        if sin_resolver:
            hallazgo("ALTA", f"{etiqueta}: marcadores sin resolver {sin_resolver[:3]}")
        if literal_none:
            hallazgo("MEDIA", f"{etiqueta}: aparece 'None' literal en el documento")
        if not sin_resolver and not literal_none:
            print(f"   ✓ {etiqueta}: documento OK")

    if activos:
        c = activos[0]
        ctx = cron("generar_informe", lambda: api_client.generar_informe_mensual(tok_admin, c["id"], MES_SIM))
        ruta = os.path.join(tempfile.gettempdir(), "sim_informe.docx")
        docx_generator.generar_informe_mensual_docx(ctx, ruta)
        verificar_docx(ruta, f"Informe docente «{c['nombre']}»")

    if directivos_puros:
        u = directivos_puros[0]
        ctx = cron("generar_gestion", lambda: api_client.generar_informe_gestion(tok_admin, u["id"], MES_SIM))
        ruta = os.path.join(tempfile.gettempdir(), "sim_gestion.docx")
        docx_generator.generar_informe_gestion_docx(ctx, ruta)
        verificar_docx(ruta, f"Informe gestión «{u['nombre']}»")

    # --- SECCIÓN E: cierre y reapertura ---------------------------------
    print(f"\n--- E. Cierre de mes (probando con {MES_CERRADO}, ya cerrado) ---")
    if algun_curso and docentes:
        # un docente intenta cargar en un mes cerrado -> debe rechazar
        tok = sesiones[algun_curso["docente_id"]]["token"]
        espera_error(f"Docente cargando en mes cerrado ({MES_CERRADO})",
                     lambda: api_client.guardar_planeacion(tok, {
                         "fecha": f"{MES_CERRADO}-07", "curso_id": algun_curso["id"],
                         "objetivo": "Objetivo. " + palabras(25), "temas_vistos": ["x"],
                         "bloques": [{"momento": "a", "minutos": 120, "observacion": palabras(25), "avance": palabras(25)}],
                         "asistencia": [{"nombre": "E", "presente": True}],
                     }, {"foto_clase": foto_payload}),
                     debe_contener="cerrado")
        # el directivo reabre y ahora sí debería poder
        cron("reabrir", lambda: api_client.reabrir_mes(tok_admin, algun_curso["id"], MES_CERRADO, True))
        try:
            r = cron("cargar tras reabrir", lambda: api_client.guardar_planeacion(tok, {
                "fecha": f"{MES_CERRADO}-07", "curso_id": algun_curso["id"],
                "objetivo": "Objetivo reapertura. " + palabras(25), "temas_vistos": ["x"],
                "bloques": [{"momento": "a", "minutos": 120, "observacion": palabras(25), "avance": palabras(25)}],
                "asistencia": [{"nombre": "E", "presente": True}],
            }, {"foto_clase": foto_payload}))
            print("   ✓ Tras reabrir, el docente pudo cargar")
            # limpiar esa planeación y volver a cerrar
            api_client.eliminar_planeacion(tok, r["id"])
            cron("cerrar de nuevo", lambda: api_client.reabrir_mes(tok_admin, algun_curso["id"], MES_CERRADO, False))
        except api_client.ApiError as e:
            hallazgo("ALTA", f"Tras reabrir el mes, el docente TODAVÍA no pudo cargar: {e}")

    # --- SECCIÓN F: limpieza --------------------------------------------
    print("\n--- F. Limpieza del mes de prueba ---")
    borradas = 0
    for c in activos:
        ses = sesiones.get(c["docente_id"])
        tok = ses["token"] if ses else tok_admin
        # informe
        try:
            api_client._call("eliminar_informe_mensual", tok_admin, c["id"], MES_SIM)
        except api_client.ApiError:
            pass
        # planeaciones del mes
        try:
            ps = api_client.obtener_planeaciones(tok, curso_id=c["id"], resumen=True)
            for p in ps:
                if str(p["fecha"])[:7] == MES_SIM:
                    api_client.eliminar_planeacion(tok, p["id"])
                    borradas += 1
        except api_client.ApiError:
            pass
        # actividades del mes
        try:
            for a in api_client.obtener_actividades(tok, c["id"], MES_SIM):
                api_client.eliminar_actividad(tok, a["id"])
        except api_client.ApiError:
            pass
    # informes de gestión + horas de gestión
    for u in directivos:
        ses = sesiones.get(u["id"])
        if not ses:
            continue
        try:
            api_client._call("eliminar_informe_mensual", tok_admin, f"gestion:{u['id']}", MES_SIM)
        except api_client.ApiError:
            pass
        try:
            for h in api_client.obtener_horas_gestion(ses["token"]):
                if str(h["fecha"])[:7] == MES_SIM:
                    api_client.eliminar_horas_gestion(ses["token"], h["id"])
        except api_client.ApiError:
            pass
    print(f"   planeaciones borradas: {borradas}")

    # verificación de limpieza
    dash_final = api_client.obtener_dashboard_directivo(tok_admin, MES_SIM)
    resto = sum(1 for d in dash_final if d["registradas"] > 0 or d["informe_entregado"])
    if resto:
        hallazgo("MEDIA", f"Quedaron {resto} cursos con datos de {MES_SIM} tras limpiar")
    else:
        print(f"   ✓ Mes {MES_SIM} limpio")

    # --- resumen ---------------------------------------------------------
    print("\n" + "=" * 60)
    print("LATENCIA (segundos por tipo de operación)")
    print("=" * 60)
    from collections import defaultdict
    agr = defaultdict(list)
    for nombre, seg in tiempos:
        clave = nombre.split()[0]
        agr[clave].append(seg)
    for clave in sorted(agr, key=lambda k: -max(agr[k])):
        v = agr[clave]
        print(f"   {clave:22} n={len(v):3}  prom={sum(v)/len(v):5.2f}  max={max(v):5.2f}")
    total = sum(seg for _, seg in tiempos)
    print(f"   TOTAL llamadas: {len(tiempos)}  tiempo acumulado: {total:.0f}s")

    print("\n" + "=" * 60)
    print(f"HALLAZGOS: {len(HALLAZGOS)}")
    print("=" * 60)
    if not HALLAZGOS:
        print("   Ninguno. Todo el flujo pasó como se esperaba.")
    for h in HALLAZGOS:
        print("   " + h)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
