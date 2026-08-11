"""Genera templates/informe_gestion.docx: el informe mensual de un
directivo sin curso (gestión pura). Ver InformesGestion.js para el contrato
de variables.

Se construye por script y no a mano para que sea reproducible: si hay que
tocar un campo, se edita acá y se vuelve a correr, en vez de pelear con
Word. Las plantillas docente (informe_mensual, planeacion) sí se hicieron
a mano a partir de los .docx reales del programa; esta no tiene un original
que copiar, así que se arma de cero.

    python scripts/construir_plantilla_gestion.py

Cuidado con los {%tr ...%}: el for y el endfor van en filas SEPARADAS de la
tabla. Inline con los datos, docxtpl tira "unknown tag 'endfor'" — nos pasó
al construir la plantilla docente.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
# Las plantillas viven en templates/ en la raíz, no en client/ (ver
# config.TEMPLATES_DIR, que en el .exe se empaqueta con --add-data).
SALIDA = RAIZ / "templates" / "informe_gestion.docx"

AZUL = RGBColor(0x1F, 0x37, 0x64)


def titulo(doc, texto, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = AZUL
    return p


def subtitulo(doc, texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = AZUL
    return p


def campo(doc, etiqueta, marcador):
    p = doc.add_paragraph()
    r = p.add_run(f"{etiqueta}: ")
    r.bold = True
    p.add_run("{{ %s }}" % marcador)


def main():
    doc = Document()

    titulo(doc, "INFORME MENSUAL DE GESTIÓN", 15)
    titulo(doc, "Generación-I", 12)

    doc.add_paragraph()
    campo(doc, "Nombre", "nombre_directivo")
    campo(doc, "Cargo", "cargo")
    campo(doc, "Cédula", "cedula")
    campo(doc, "Periodo evaluado", "periodo_evaluado")

    doc.add_paragraph()
    subtitulo(doc, "1. Desarrollo de la gestión del mes")
    campo(doc, "Objetivos y metas", "gestion_objetivos")
    campo(doc, "Principales logros y entregables", "gestion_logros")
    campo(doc, "Novedades, obstáculos o riesgos", "gestion_novedades")
    campo(doc, "Estrategias y acciones correctivas", "gestion_estrategias")
    campo(doc, "Pendientes para el próximo mes", "gestion_pendientes")

    doc.add_paragraph()
    subtitulo(doc, "2. Detalle de horas de gestión")

    # Tabla: encabezado + fila-for + fila-datos + fila-endfor. El for y el
    # endfor viven en su propia fila, sin datos, o docxtpl rompe.
    tabla = doc.add_table(rows=1, cols=5)
    tabla.style = "Table Grid"
    encabezados = ["Actividad / tarea", "Semana", "Horas", "Entregable", "Soporte"]
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        run = celda.paragraphs[0].add_run(texto)
        run.bold = True

    fila_for = tabla.add_row().cells
    fila_for[0].paragraphs[0].add_run("{%tr for h in horas_gestion %}")

    fila_datos = tabla.add_row().cells
    fila_datos[0].paragraphs[0].add_run("{{ h.actividad }}")
    fila_datos[1].paragraphs[0].add_run("{{ h.nro_semana }}")
    fila_datos[2].paragraphs[0].add_run("{{ h.horas_sede }}")
    fila_datos[3].paragraphs[0].add_run("{{ h.entregable }}")
    fila_datos[4].paragraphs[0].add_run("{{ h.link_soporte }}")

    fila_endfor = tabla.add_row().cells
    fila_endfor[0].paragraphs[0].add_run("{%tr endfor %}")

    p = doc.add_paragraph()
    r = p.add_run("Total de horas de gestión: ")
    r.bold = True
    p.add_run("{{ total_horas_gestion }}")

    doc.add_paragraph()
    subtitulo(doc, "3. Cuenta de cobro")
    campo(doc, "Mes a cobrar", "mes_a_cobrar")
    campo(doc, "Periodo", "periodo_cobro")
    campo(doc, "Horas totales", "horas_totales")
    campo(doc, "Valor a cobrar", "valor_en_numeros")
    campo(doc, "Son", "valor_en_letras")
    campo(doc, "Número de cuenta", "numero_cuenta")
    campo(doc, "Tipo de cuenta", "tipo_cuenta")
    campo(doc, "Entidad bancaria", "entidad_bancaria")
    campo(doc, "Fecha de emisión", "fecha_emision")

    doc.add_paragraph()
    doc.add_paragraph()
    # La firma es una imagen que docxtpl inserta con InlineImage; el
    # cliente reemplaza {{ firma }} por la imagen o por "" si no hay.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("{{ firma }}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("{{ nombre_directivo }}").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Firma del directivo")

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(SALIDA))
    print(f"Plantilla generada en {SALIDA}")


if __name__ == "__main__":
    main()
