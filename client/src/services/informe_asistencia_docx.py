"""Informe mensual de asistencia (formato nuevo, una página por curso).

A diferencia de las demás planeaciones/informes, este .docx NO sale de una
plantilla docxtpl sino que se arma directo con python-docx: las columnas
cambian con la cantidad de clases del mes (3, 4, 5, 6 o más) y una plantilla
fija tenía que dejar seis columnas de tope. Ver Asistencia.js para el
contrato de datos de cada curso.

Pensado para imprimirse también en blanco y negro:
- cada estado se distingue por su símbolo Y su palabra (✓ Asistió, ✗ Faltó,
  — Retirado), nunca solo por el color;
- los rellenos son tenues (en gris se leen como un sombreado suave) y el
  texto siempre va oscuro sobre fondo claro;
- el encabezado de la tabla es gris claro con texto negro, no un bloque
  negro con letra blanca (que en fotocopiadoras y láser económicas se tapa).
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from config import TEMPLATES_DIR

LOGO = TEMPLATES_DIR / "assets" / "logo_generacion_i_completo.png"

FUENTE = "Arial"
FUENTE_SIMBOLOS = "Segoe UI Symbol"  # Arial no trae ✓ ✗; Word lo cambiaría por otra de todos modos

TINTA = "111111"
APAGADO = "4B5563"
GRIS_LINEA = "9CA3AF"
GRIS_ENCABEZADO = "E5E7EB"
GRIS_LEYENDA = "F3F4F6"
# Asistió muy claro y Faltó más marcado: en color son verde y rojo suaves, y en
# blanco y negro quedan como dos grises distintos (~96 % y ~85 % de claridad).
ASISTIO_FONDO, ASISTIO_TEXTO = "EEF7F0", "14532D"
FALTO_FONDO, FALTO_TEXTO = "F6D0D0", "8B1A1A"

ANCHO_UTIL_CM = 18.2  # A4 vertical con márgenes de 1,4 cm
MAX_CLASES_CON_TEXTO = 6  # más columnas que esto y la marca pasa a solo símbolo

_SIMBOLO = {"asistio": "✓", "falto": "✗", "retirado": "—"}
_PALABRA = {"asistio": "Asistió", "falto": "Faltó", "retirado": "Retirado"}

# Orden que exige el esquema de Word para los hijos de cada propiedad: si
# un elemento queda fuera de lugar, Word puede negarse a abrir el archivo.
_ORDEN_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
               "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
_ORDEN_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
                "shd", "tblLayout", "tblCellMar", "tblLook"]
_ORDEN_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
              "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
              "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
              "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
              "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"]


# ---------------------------------------------------------------- XML

def _insertar(padre, hijo, orden):
    """Mete `hijo` en `padre` respetando el orden del esquema, reemplazando
    uno del mismo tipo si ya estaba."""
    nombre = hijo.tag.split("}")[1]
    for existente in padre.findall(qn(f"w:{nombre}")):
        padre.remove(existente)
    posicion = orden.index(nombre)
    for i, actual in enumerate(list(padre)):
        otro = actual.tag.split("}")[1]
        if otro in orden and orden.index(otro) > posicion:
            padre.insert(i, hijo)
            return
    padre.append(hijo)


def _borde_el(lado, valor="single", tam=4, color=GRIS_LINEA, espacio=0):
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), valor)
    if valor != "nil":
        el.set(qn("w:sz"), str(tam))
        el.set(qn("w:space"), str(espacio))
        el.set(qn("w:color"), color)
    return el


def _bordes_celda(celda, **lados):
    """lados: top/left/bottom/right = None (sin borde) o (val, sz, color)."""
    tcPr = celda._tc.get_or_add_tcPr()
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        spec = lados.get(lado)
        bordes.append(_borde_el(lado, "nil") if spec is None else _borde_el(lado, *spec))
    _insertar(tcPr, bordes, _ORDEN_TCPR)


def _rellenar_celda(celda, color_hex):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    _insertar(tcPr, shd, _ORDEN_TCPR)


def _margenes_celda(celda, arriba=40, abajo=40, izq=70, der=70):
    tcPr = celda._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insertar(tcPr, mar, _ORDEN_TCPR)


def _centrar_vertical(celda):
    tcPr = celda._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    _insertar(tcPr, va, _ORDEN_TCPR)


def _borde_parrafo(parrafo, **lados):
    """lados: top/left/bottom/right = (val, sz, color, space)."""
    pPr = parrafo._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        if lado not in lados:
            continue
        val, tam, color, espacio = lados[lado]
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(tam))
        el.set(qn("w:space"), str(espacio))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    _insertar(pPr, pbdr, _ORDEN_PPR)


def _rellenar_parrafo(parrafo, color_hex):
    pPr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    _insertar(pPr, shd, _ORDEN_PPR)


def _tabla(doc, anchos_cm):
    """Tabla de tamaño fijo y sin bordes propios (cada celda pone los suyos)."""
    tabla = doc.add_table(rows=1, cols=len(anchos_cm))
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    ancho = OxmlElement("w:tblW")
    ancho.set(qn("w:w"), str(int(sum(anchos_cm) * 567)))
    ancho.set(qn("w:type"), "dxa")
    _insertar(tblPr, ancho, _ORDEN_TBLPR)
    # Word corre la tabla hacia la izquierda el margen interno por defecto
    # de sus celdas: en cero, el borde de la tabla queda justo en el margen
    # de la página, alineado con las reglas y los títulos.
    sangria = OxmlElement("w:tblInd")
    sangria.set(qn("w:w"), "0")
    sangria.set(qn("w:type"), "dxa")
    _insertar(tblPr, sangria, _ORDEN_TBLPR)
    mar = OxmlElement("w:tblCellMar")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insertar(tblPr, mar, _ORDEN_TBLPR)
    for col, a in zip(tabla.columns, anchos_cm):
        col.width = Cm(a)
    for celda, a in zip(tabla.rows[0].cells, anchos_cm):
        celda.width = Cm(a)
    return tabla


def _fila(tabla, anchos_cm, alto_cm=None, repetir=False):
    fila = tabla.add_row()
    for celda, a in zip(fila.cells, anchos_cm):
        celda.width = Cm(a)
    _configurar_fila(fila, alto_cm, repetir)
    return fila


def _configurar_fila(fila, alto_cm=None, repetir=False):
    if alto_cm:
        fila.height = Cm(alto_cm)
        fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if repetir:
        trPr.append(OxmlElement("w:tblHeader"))


# ------------------------------------------------------------- texto

def _fuente(run, nombre):
    run.font.name = nombre
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    for attr in ("w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), nombre)


def _run(parrafo, texto, tam=9, negrita=False, cursiva=False, color=TINTA, fuente=FUENTE):
    r = parrafo.add_run(texto)
    r.font.size = Pt(tam)
    r.font.bold = negrita
    r.font.italic = cursiva
    r.font.color.rgb = RGBColor.from_string(color)
    _fuente(r, fuente)
    return r


def _chip(parrafo, texto, tam=7, color=TINTA, fondo=None, borde=GRIS_LINEA, negrita=True):
    """Texto con marco (y relleno opcional) alrededor, como una etiqueta.
    El marco es lo que la distingue en blanco y negro cuando el relleno
    sale casi blanco."""
    r = _run(parrafo, f" {texto} ", tam=tam, negrita=negrita, color=color)
    rPr = r._r.get_or_add_rPr()
    bdr = OxmlElement("w:bdr")
    bdr.set(qn("w:val"), "single")
    bdr.set(qn("w:sz"), "4")
    bdr.set(qn("w:space"), "0")
    bdr.set(qn("w:color"), borde)
    rPr.append(bdr)
    if fondo:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fondo)
        rPr.append(shd)
    return r


def _parrafo(celda_o_doc, alineacion=None, antes=0, despues=0, mantener=False, primero=False):
    if primero and hasattr(celda_o_doc, "paragraphs"):
        p = celda_o_doc.paragraphs[0]
    else:
        p = celda_o_doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(antes)
    f.space_after = Pt(despues)
    f.line_spacing = 1.0
    if alineacion is not None:
        p.alignment = alineacion
    if mantener:
        f.keep_with_next = True
    return p


def _separador(doc, pt=4):
    p = _parrafo(doc)
    _run(p, "", tam=1)
    p.paragraph_format.line_spacing = Pt(pt)
    return p


# --------------------------------------------------------- normalización

def _estado(marca: str) -> str:
    """Acepta lo que mande el backend: el nuevo ("Retirado") y el anterior
    ("—"), para que un cliente nuevo siga sirviendo con un backend sin
    actualizar."""
    if marca == "Asistió":
        return "asistio"
    if marca in ("Retirado", "—"):
        return "retirado"
    return "falto"


def _estudiante(est: dict, n_clases: int) -> dict:
    nombre = str(est.get("nombre", "")).strip()
    retirado = bool(est.get("retirado"))
    if nombre.lower().endswith("(retirado)"):
        nombre = nombre[: -len("(retirado)")].strip()
        retirado = True
    marcas = [_estado(m) for m in (est.get("marcas") or [])]
    marcas += ["falto"] * (n_clases - len(marcas))
    return {
        "nombre": nombre,
        "retirado": retirado,
        "marcas": marcas[:n_clases],
        "asistio": est.get("total_asistio", marcas.count("asistio")),
        "falto": est.get("total_falto", marcas.count("falto")),
    }


# ---------------------------------------------------------------- logo

_logo_bytes: bytes | None = None


def _logo() -> io.BytesIO | None:
    global _logo_bytes
    if _logo_bytes is None:
        if not Path(LOGO).exists():
            return None
        _logo_bytes = Path(LOGO).read_bytes()
    return io.BytesIO(_logo_bytes)


# ------------------------------------------------------------- bloques

def _encabezado(doc, mes_nombre, anio, fecha_emision):
    anchos = [4.4, 4.6, 9.2]
    t = _tabla(doc, anchos)
    c_logo, _, c_titulo = t.rows[0].cells
    for c in t.rows[0].cells:
        _bordes_celda(c)
        _margenes_celda(c, 0, 0, 0, 0)
        _centrar_vertical(c)

    p = _parrafo(c_logo, primero=True)
    logo = _logo()
    if logo:
        p.add_run().add_picture(logo, width=Cm(3.9))

    p = _parrafo(c_titulo, WD_ALIGN_PARAGRAPH.RIGHT, primero=True, despues=5)
    _run(p, "INFORME MENSUAL DE ASISTENCIA", tam=11, negrita=True)
    _borde_parrafo(p, top=("single", 12, TINTA, 4), bottom=("single", 12, TINTA, 4),
                   left=("single", 12, TINTA, 6), right=("single", 12, TINTA, 6))
    p = _parrafo(c_titulo, WD_ALIGN_PARAGRAPH.RIGHT, antes=3)
    _run(p, f"Periodo: {mes_nombre} {anio}", tam=9, negrita=True)
    _run(p, f"   |   Emisión: {fecha_emision}", tam=9, color=APAGADO)

    regla = _parrafo(doc, antes=4, despues=6)
    _run(regla, "", tam=1)
    _borde_parrafo(regla, bottom=("single", 18, TINTA, 1))


def _tarjeta_curso(doc, curso, estudiantes):
    total_clases = len(curso.get("clases") or [])
    retirados = sum(1 for e in estudiantes if e["retirado"])
    porcentaje = curso.get("porcentaje_asistencia", 0)

    anchos = [8.3, 0.25, 3.15, 0.25, 3.1, 0.25, 2.9]
    t = _tabla(doc, anchos)
    celdas = t.rows[0].cells
    _configurar_fila(t.rows[0], 1.9)
    for i, c in enumerate(celdas):
        _margenes_celda(c, 60, 60, 110, 110)
        _centrar_vertical(c)
        if i % 2 == 1:
            _bordes_celda(c)
        else:
            _bordes_celda(c, top=("single", 6, TINTA, 0), bottom=("single", 6, TINTA, 0),
                          right=("single", 6, TINTA, 0), left=("single", 6, TINTA, 0))
    _bordes_celda(celdas[0], top=("single", 6, TINTA, 0), bottom=("single", 6, TINTA, 0),
                  right=("single", 6, TINTA, 0), left=("single", 30, TINTA, 0))

    p = _parrafo(celdas[0], primero=True, despues=2)
    _run(p, curso.get("nombre", ""), tam=13, negrita=True)
    p = _parrafo(celdas[0], despues=2)
    if curso.get("nucleo"):
        _chip(p, f"NÚCLEO: {curso['nucleo']}".upper(), tam=6.5)
    p = _parrafo(celdas[0])
    _run(p, "Docente titular: ", tam=9, color=APAGADO)
    _run(p, curso.get("docente", "") or "—", tam=9, negrita=True)

    def estadistica(celda, valor, etiqueta, nota=None):
        p = _parrafo(celda, WD_ALIGN_PARAGRAPH.CENTER, primero=True)
        _run(p, str(valor), tam=20, negrita=True)
        p = _parrafo(celda, WD_ALIGN_PARAGRAPH.CENTER)
        _run(p, etiqueta, tam=6.5, negrita=True, color=APAGADO)
        if nota:
            p = _parrafo(celda, WD_ALIGN_PARAGRAPH.CENTER)
            _run(p, nota, tam=6.5, color=APAGADO)

    estadistica(celdas[2], total_clases, "CLASES DICTADAS")
    nota = None
    if retirados:
        nota = f"({retirados} retirado{'s' if retirados != 1 else ''})"
    estadistica(celdas[4], len(estudiantes), "ESTUDIANTES", nota)
    estadistica(celdas[6], f"{porcentaje}%", "ASISTENCIA PROMEDIO")


def _leyenda(doc):
    p = _parrafo(doc, antes=8, despues=5)
    p.paragraph_format.left_indent = Cm(0.3)  # para que la barra izquierda caiga en el margen
    _rellenar_parrafo(p, GRIS_LEYENDA)
    _borde_parrafo(p, left=("single", 30, TINTA, 4))
    _run(p, "  REGISTRO DETALLADO POR SESIÓN      ", tam=8.5, negrita=True)
    for estado, fondo, texto, aclaracion in (
        ("asistio", ASISTIO_FONDO, ASISTIO_TEXTO, "Presente"),
        ("falto", FALTO_FONDO, FALTO_TEXTO, "Inasistencia"),
        ("retirado", None, APAGADO, "No era del curso en esa clase"),
    ):
        _rellenar_run_simbolo(p, estado, fondo, texto)
        _run(p, f" {aclaracion}    ", tam=7.5, color=APAGADO)


def _rellenar_run_simbolo(parrafo, estado, fondo, color):
    """El chip de la leyenda: símbolo (en su fuente) + palabra, con marco."""
    for texto, fuente in ((f" {_SIMBOLO[estado]}", FUENTE_SIMBOLOS), (f" {_PALABRA[estado]} ", FUENTE)):
        r = _run(parrafo, texto, tam=7.5, negrita=True, cursiva=estado == "retirado",
                 color=color, fuente=fuente)
        rPr = r._r.get_or_add_rPr()
        bdr = OxmlElement("w:bdr")
        bdr.set(qn("w:val"), "single")
        bdr.set(qn("w:sz"), "4")
        bdr.set(qn("w:space"), "0")
        bdr.set(qn("w:color"), GRIS_LINEA)
        rPr.append(bdr)
        if fondo:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fondo)
            rPr.append(shd)


def _celda_marca(celda, estado, compacto):
    fondo, color = {
        "asistio": (ASISTIO_FONDO, ASISTIO_TEXTO),
        "falto": (FALTO_FONDO, FALTO_TEXTO),
        "retirado": (None, APAGADO),
    }[estado]
    if fondo:
        _rellenar_celda(celda, fondo)
    _centrar_vertical(celda)
    p = _parrafo(celda, WD_ALIGN_PARAGRAPH.CENTER, primero=True)
    cursiva = estado == "retirado"
    if compacto:
        _run(p, _SIMBOLO[estado], tam=9, negrita=True, color=color, fuente=FUENTE_SIMBOLOS)
        return
    _run(p, _SIMBOLO[estado] + " ", tam=7.5, negrita=True, cursiva=cursiva, color=color,
         fuente=FUENTE_SIMBOLOS)
    _run(p, _PALABRA[estado], tam=7.5, negrita=True, cursiva=cursiva, color=color)


def _tabla_asistencia(doc, clases, estudiantes):
    n = len(clases)
    compacto = n > MAX_CLASES_CON_TEXTO
    w_num, w_tot = 0.7, 1.15
    w_clase = 2.35 if n <= 3 else (1.85 if not compacto else max(0.95, min(1.6, (ANCHO_UTIL_CM - 2 * w_tot - w_num - 4.6) / max(n, 1))))
    w_nombre = ANCHO_UTIL_CM - w_num - 2 * w_tot - w_clase * n
    anchos = [w_num, w_nombre] + [w_clase] * n + [w_tot, w_tot]

    t = _tabla(doc, anchos)
    encabezados = ["#", "ESTUDIANTE"] + [None] * n + ["ASIST.", "FALT."]
    fila0 = t.rows[0]
    _configurar_fila(fila0, 0.95, repetir=True)
    for i, celda in enumerate(fila0.cells):
        _rellenar_celda(celda, GRIS_ENCABEZADO)
        _bordes_celda(celda, top=("single", 12, TINTA, 0), bottom=("single", 12, TINTA, 0),
                      left=("single", 4, GRIS_LINEA, 0), right=("single", 4, GRIS_LINEA, 0))
        _margenes_celda(celda, 50, 50, 50, 50)
        _centrar_vertical(celda)
        alin = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
        p = _parrafo(celda, alin, primero=True)
        if 2 <= i < 2 + n:
            clase = clases[i - 2]
            nro = str(clase.get("nro", "")).upper()
            fecha = str(clase.get("fecha", ""))
            _run(p, nro.replace("CLASE ", "C") if compacto else nro, tam=7, negrita=True)
            p2 = _parrafo(celda, WD_ALIGN_PARAGRAPH.CENTER)
            _run(p2, fecha[:5] if compacto else fecha, tam=6.5, color=APAGADO)
        else:
            _run(p, encabezados[i], tam=7.5, negrita=True)

    for pos, est in enumerate(estudiantes, start=1):
        fila = _fila(t, anchos, 0.62)
        ultima = pos == len(estudiantes)
        for i, celda in enumerate(fila.cells):
            _margenes_celda(celda, 35, 35, 60, 60)
            _centrar_vertical(celda)
            _bordes_celda(
                celda,
                top=("single", 4, GRIS_LINEA, 0),
                bottom=("single", 8 if ultima else 4, TINTA if ultima else GRIS_LINEA, 0),
                left=("single", 4, GRIS_LINEA, 0), right=("single", 4, GRIS_LINEA, 0),
            )
            if ultima:
                for p in celda.paragraphs:
                    p.paragraph_format.keep_with_next = True

        c = fila.cells
        p = _parrafo(c[0], WD_ALIGN_PARAGRAPH.CENTER, primero=True, mantener=ultima)
        _run(p, str(pos), tam=7.5, color=APAGADO)

        p = _parrafo(c[1], primero=True, mantener=ultima)
        _run(p, est["nombre"], tam=8.5, negrita=True)
        if est["retirado"]:
            _run(p, "  ", tam=6)
            _chip(p, "RETIRADO", tam=6, color=APAGADO, fondo=GRIS_ENCABEZADO)

        for j, estado in enumerate(est["marcas"]):
            _celda_marca(c[2 + j], estado, compacto)
            for p in c[2 + j].paragraphs:
                p.paragraph_format.keep_with_next = ultima

        p = _parrafo(c[2 + n], WD_ALIGN_PARAGRAPH.CENTER, primero=True, mantener=ultima)
        _run(p, str(est["asistio"]), tam=9, negrita=True)
        p = _parrafo(c[3 + n], WD_ALIGN_PARAGRAPH.CENTER, primero=True, mantener=ultima)
        _run(p, str(est["falto"]), tam=9, negrita=True)


def _firmas(doc, curso):
    _separador(doc, 26)
    anchos = [8.4, 1.4, 8.4]
    t = _tabla(doc, anchos)
    _configurar_fila(t.rows[0], 1.3)
    fila_nombre = _fila(t, anchos)
    fila_cargo = _fila(t, anchos)

    coordinadora = curso.get("coordinadora") or ""
    columnas = (
        (0, curso.get("docente", "") or "—", "DOCENTE TITULAR", False),
        (2, coordinadora or "(sin coordinación asignada para este curso)",
         "COORDINACIÓN", not coordinadora),
    )
    for fila in t.rows:
        for c in fila.cells:
            _margenes_celda(c, 20, 20, 60, 60)
            _bordes_celda(c)
    for idx, nombre, cargo, vacio in columnas:
        _bordes_celda(t.rows[0].cells[idx], bottom=("single", 8, TINTA, 0))
        p = _parrafo(fila_nombre.cells[idx], WD_ALIGN_PARAGRAPH.CENTER, primero=True, antes=2, mantener=True)
        _run(p, nombre, tam=9, negrita=not vacio, cursiva=vacio, color=APAGADO if vacio else TINTA)
        p = _parrafo(fila_cargo.cells[idx], WD_ALIGN_PARAGRAPH.CENTER, primero=True)
        _run(p, cargo, tam=6.5, negrita=True, color=APAGADO)
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = True


def _pie(seccion):
    pie = seccion.footer
    pie.is_linked_to_previous = False
    p = pie.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(Cm(ANCHO_UTIL_CM), WD_TAB_ALIGNMENT.RIGHT)
    _borde_parrafo(p, top=("single", 6, GRIS_LINEA, 4))
    _run(p, "Programa CTeI Generación-i • Alcaldía de San Pedro de los Milagros, Antioquia",
         tam=7.5, color=APAGADO)
    _run(p, "\tPágina ", tam=7.5, color=APAGADO)
    _campo(p, "PAGE")
    _run(p, " de ", tam=7.5, color=APAGADO)
    _campo(p, "NUMPAGES")


def _campo(parrafo, instruccion):
    """Campo de Word (número de página / total). Se actualiza solo al abrir
    o al convertir a PDF."""
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), (None, "1"), ("end", None)):
        r = _run(parrafo, "", tam=7.5, color=APAGADO)
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r._r.append(fc)
        elif texto == instruccion:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instruccion} "
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = texto
            r._r.append(t)


def _configurar_pagina(seccion):
    seccion.orientation = WD_ORIENT.PORTRAIT
    seccion.page_width, seccion.page_height = Cm(21.0), Cm(29.7)
    seccion.left_margin = seccion.right_margin = Cm(1.4)
    seccion.top_margin, seccion.bottom_margin = Cm(1.2), Cm(1.5)
    seccion.footer_distance = Cm(0.7)


def _hoja_curso(doc, mes_nombre, anio, fecha_emision, curso):
    clases = curso.get("clases") or []
    estudiantes = [_estudiante(e, len(clases)) for e in (curso.get("estudiantes") or [])]
    _encabezado(doc, mes_nombre, anio, fecha_emision)
    _tarjeta_curso(doc, curso, estudiantes)
    _leyenda(doc)
    if not estudiantes:
        p = _parrafo(doc, antes=10)
        _run(p, "Este curso no tiene estudiantes registrados.", tam=9, cursiva=True, color=APAGADO)
    else:
        _tabla_asistencia(doc, clases, estudiantes)
    _firmas(doc, curso)


def _tabla_simple(doc, encabezados, filas, anchos):
    t = _tabla(doc, anchos)
    for celda, texto in zip(t.rows[0].cells, encabezados):
        _rellenar_celda(celda, GRIS_ENCABEZADO)
        _bordes_celda(celda, top=("single", 12, TINTA, 0), bottom=("single", 12, TINTA, 0))
        _margenes_celda(celda, 50, 50, 70, 70)
        _run(_parrafo(celda, primero=True), texto, tam=8, negrita=True)
    for valores in filas:
        fila = _fila(t, anchos, 0.6)
        for celda, valor in zip(fila.cells, valores):
            _margenes_celda(celda, 35, 35, 70, 70)
            _bordes_celda(celda, bottom=("single", 4, GRIS_LINEA, 0))
            _run(_parrafo(celda, primero=True), str(valor), tam=8.5)


def _anexos(doc, contexto):
    """Lo que no es de un curso puntual: los activos sin clases del mes y (si
    un backend anterior todavía lo manda) los nombres que no se pudieron
    cruzar con la lista de inscritos."""
    sin_clases = contexto.get("cursos_sin_clases") or []
    no_identificados = contexto.get("no_identificados") or []
    if not sin_clases and not no_identificados:
        return
    doc.add_section(WD_SECTION.NEW_PAGE)
    p = _parrafo(doc, despues=6)
    _run(p, "OBSERVACIONES DEL PERIODO", tam=12, negrita=True)
    _borde_parrafo(p, bottom=("single", 12, TINTA, 2))
    if sin_clases:
        p = _parrafo(doc, antes=8, despues=4)
        _run(p, f"Cursos activos sin ninguna clase cargada este mes ({len(sin_clases)})", tam=9.5, negrita=True)
        _tabla_simple(doc, ["CURSO", "DOCENTE"],
                      [(c.get("nombre", ""), c.get("docente", "")) for c in sin_clases], [9.0, 9.2])
    if no_identificados:
        p = _parrafo(doc, antes=12, despues=4)
        _run(p, "Nombres marcados presentes que no coinciden con la lista de inscritos "
                f"({len(no_identificados)})", tam=9.5, negrita=True)
        _tabla_simple(doc, ["NOMBRE", "CURSO", "CLASE", "FECHA"],
                      [(n.get("nombre", ""), n.get("curso", ""), n.get("clase", ""), n.get("fecha", ""))
                       for n in no_identificados], [6.2, 6.6, 2.7, 2.7])


def generar(contexto: dict, ruta_salida: str | Path, con_anexos: bool = True) -> Path:
    """contexto: mes_nombre, anio, fecha_emision, cursos[] (y opcionalmente
    cursos_sin_clases[] / no_identificados[]). Una página por curso."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    _configurar_pagina(doc.sections[0])
    _pie(doc.sections[0])
    doc.core_properties.title = f"Informe de asistencia — {contexto.get('mes_nombre', '')} {contexto.get('anio', '')}"

    cursos = contexto.get("cursos") or []
    for i, curso in enumerate(cursos):
        if i > 0:
            doc.add_section(WD_SECTION.NEW_PAGE)
        _hoja_curso(doc, contexto.get("mes_nombre", ""), contexto.get("anio", ""),
                    contexto.get("fecha_emision", ""), curso)
    if con_anexos:
        _anexos(doc, contexto)

    ruta_salida = Path(ruta_salida)
    doc.save(str(ruta_salida))
    return ruta_salida
