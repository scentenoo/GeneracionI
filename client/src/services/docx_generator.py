"""Relleno de las plantillas docxtpl (templates/planeacion_individual.docx y
templates/informe_mensual.docx). Ver esos archivos para el contrato exacto
de variables — quedaron documentados ahí mismo al construirlos.

Esto es lógica de Samir + Claude Code (spec sección 8): las pantallas de
los niños solo deberían llamar a estas dos funciones con los datos ya
armados, nunca tocar docxtpl directamente.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docxtpl import DocxTemplate, InlineImage
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from config import TEMPLATES_DIR
from services import informe_asistencia_docx

PLANEACION_TEMPLATE = TEMPLATES_DIR / "planeacion_individual.docx"
INFORME_TEMPLATE = TEMPLATES_DIR / "informe_mensual.docx"
INFORME_GESTION_TEMPLATE = TEMPLATES_DIR / "informe_gestion.docx"
CERTIFICADO_PAGO_TEMPLATE = TEMPLATES_DIR / "certificado_pago.docx"

_IMG_WIDTH_GRANDE_MM = 90
_IMG_WIDTH_CHICA_MM = 55

# Ancho de cada foto de clase según cuántas haya (1 a 3), para que entren
# en una sola fila del documento sin pasarse del margen de la página.
_ANCHO_FOTO_CLASE_MM = {1: _IMG_WIDTH_GRANDE_MM, 2: 75, 3: 55}

# Cómo se lee cada acción del historial en la hoja final del documento.
_ACCIONES = {
    "entregado": "Entregado",
    "actualizado": "Actualizado",
    "reenviado": "Reenviado corregido",
    "devuelto": "Devuelto para corregir",
    "aprobado": "Aprobado",
}


def _agregar_bordes_tabla(tabla) -> None:
    """Bordes finos en toda la tabla, puestos a mano en el XML en vez de con
    `tabla.style = "Table Grid"` — ese nombre de estilo solo existe si la
    plantilla .docx ya lo usó alguna vez (Word los agrega recién al primer
    uso); informe_mensual.docx no lo tenía y eso tiraba
    KeyError: no style with name 'Table Grid' justo al final, después de
    armar todo el resto del documento."""
    tblPr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        bordes.append(el)
    tblPr.append(bordes)


_TITULO_HISTORIAL = "Historial de revisión"


def _anexar_historial_a_doc(doc: Document, historial: list | None) -> None:
    """Agrega al final del documento (ya abierto) la hoja de auditoría
    —quién lo entregó, quién lo devolvió y por qué, quién lo aprobó y
    cuándo—, pedida por dirección para que quede en TODOS los documentos
    (no solo los que se devolvieron alguna vez): un registro completo de
    principio a fin, aunque se haya aprobado a la primera."""
    if not historial:
        return

    doc.add_page_break()

    titulo = doc.add_paragraph()
    run = titulo.add_run(_TITULO_HISTORIAL)
    run.bold = True
    run.font.size = Pt(14)

    sub = doc.add_paragraph()
    r = sub.add_run(
        "Registro de entregas, devoluciones y aprobación de este documento."
    )
    r.italic = True
    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    tabla = doc.add_table(rows=1, cols=4)
    _agregar_bordes_tabla(tabla)
    encabezados = ["Fecha", "Acción", "Responsable", "Motivo"]
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        p = celda.paragraphs[0]
        run = p.add_run(texto)
        run.bold = True

    for e in historial:
        fila = tabla.add_row().cells
        fecha = str(e.get("fecha", "")).replace("T", " ")[:16]
        fila[0].text = fecha
        fila[1].text = _ACCIONES.get(e.get("accion", ""), e.get("accion", ""))
        fila[2].text = str(e.get("autor", ""))
        fila[3].text = str(e.get("motivo", ""))


def _quitar_historial_previo(doc: Document) -> None:
    """Si el documento ya traía una hoja de historial de una versión
    anterior, la saca antes de agregar la actualizada — para que
    `actualizar_historial_docx` reemplace la hoja en vez de apilar una
    nueva encima de la vieja cada vez que se aprueba o se devuelve.

    Busca el párrafo "Historial de revisión" y borra, en orden: el salto
    de página que `_anexar_historial_a_doc` pone justo antes, y todo lo que
    viene después (subtítulo y tabla) sin tocar el `sectPr` final del
    cuerpo — ese define el tamaño de página y hay que dejarlo donde está."""
    body = doc.element.body
    for p in doc.paragraphs:
        if p.text != _TITULO_HISTORIAL:
            continue
        encabezado_el = p._p
        anterior = encabezado_el.getprevious()
        if anterior is not None:
            body.remove(anterior)
        siguiente = encabezado_el.getnext()
        while siguiente is not None and siguiente.tag != qn("w:sectPr"):
            actual, siguiente = siguiente, siguiente.getnext()
            body.remove(actual)
        body.remove(encabezado_el)
        return


def _anexar_historial(ruta_salida: Path, historial: list | None) -> None:
    """Como `_anexar_historial_a_doc`, pero abre y guarda el .docx por su
    ruta — lo que usan `generar_*_docx` después de que docxtpl ya
    renderizó el resto del documento."""
    if not historial:
        return
    doc = Document(str(ruta_salida))
    _anexar_historial_a_doc(doc, historial)
    doc.save(str(ruta_salida))


# Cuántas líneas de texto entran en la hoja de historial sin desbordar a una
# segunda página. Una hoja tiene unas 40; se deja un colchón grande porque
# esto se ESTIMA a ojo (sin Word no hay cómo medir), y equivocarse hacia
# arriba solo cuesta un recuento de páginas de más — hacia abajo, dejar un
# "2 de 1" en el pie.
_LINEAS_MAX_HOJA_HISTORIAL = 28
# Cada columna de la tabla mide ~1/4 del ancho útil: de ahí ~14 caracteres
# por línea, contando con que una palabra larga corta antes.
_CARACTERES_POR_LINEA = 14


def _lineas_de_fila(celdas) -> int:
    """Líneas que ocupa una fila de la tabla de historial: la de su celda
    más larga, más una de aire entre filas."""
    return 1 + max(
        max(1, -(-len(str(c)) // _CARACTERES_POR_LINEA)) for c in celdas
    )


def _lineas_historial(historial: list) -> int:
    """Estimación de cuántas líneas ocupa la hoja recién armada por
    `_anexar_historial_a_doc` (encabezado de la tabla incluido)."""
    lineas = _lineas_de_fila(["Fecha", "Acción", "Responsable", "Motivo"])
    for e in historial:
        lineas += _lineas_de_fila((
            str(e.get("fecha", "")).replace("T", " ")[:16],
            _ACCIONES.get(e.get("accion", ""), e.get("accion", "")),
            e.get("autor", ""),
            e.get("motivo", ""),
        ))
    return lineas


def _lineas_historial_en_doc(doc: Document) -> int | None:
    """Lo mismo pero sobre la hoja que el documento ya trae. None si no
    tiene hoja de historial (documento anterior a que existiera)."""
    if not any(p.text == _TITULO_HISTORIAL for p in doc.paragraphs) or not doc.tables:
        return None
    # La del historial es siempre la última tabla del cuerpo: se agrega al final.
    return sum(_lineas_de_fila([c.text for c in fila.cells]) for fila in doc.tables[-1].rows)


def actualizar_historial_docx(entrada, historial: list, ruta_salida: str | Path) -> bool:
    """Reemplaza SOLO la hoja de "Historial de revisión" de un .docx que
    ya existe, sin volver a renderizar el resto del documento — para que
    aprobar o devolver una planeación (o un informe) sea rápido: no hace
    falta re-descargar fotos ni rellenar la plantilla de nuevo, solo
    cambia la última página.

    `entrada`: ruta o objeto tipo archivo (p.ej. io.BytesIO) con el .docx
    ya archivado. `historial` no puede venir vacío — si el documento no
    tiene por qué llevar hoja de historial, no hay nada que actualizar.

    Devuelve True si hay que volver a contar las páginas del documento
    (ver pdf_converter.recalcular_campos) porque el total pudo haber
    cambiado; False si se puede asegurar que la hoja sigue ocupando una
    sola página, antes y después, y el "Página X de Y" queda tal cual —
    que es lo normal, y evita arrancar Word (~6 s) por cada revisión."""
    doc = Document(entrada)
    antes = _lineas_historial_en_doc(doc)
    _quitar_historial_previo(doc)
    _anexar_historial_a_doc(doc, historial)
    doc.save(str(Path(ruta_salida)))
    return (
        antes is None
        or antes > _LINEAS_MAX_HOJA_HISTORIAL
        or _lineas_historial(historial) > _LINEAS_MAX_HOJA_HISTORIAL
    )


def _imagen_desde_base64(tpl: DocxTemplate, base64_str: str | None, ancho_mm: int) -> InlineImage | str:
    """Si no hay imagen (ej. todavía no se subió la firma), devuelve '' en
    vez de None — con None, docxtpl renderiza el texto literal "None"."""
    if not base64_str:
        return ""
    data = base64.b64decode(base64_str)
    return InlineImage(tpl, io.BytesIO(data), width=Mm(ancho_mm))


def generar_planeacion_docx(
    contexto: dict, fotos_clase_paths: list[str], ruta_salida: str | Path
) -> Path:
    """contexto: fecha, grupo, objetivo, temas_vistos[], los tres momentos
    (momento_*_min/texto), observaciones (evaluación de la clase),
    asistencia[]. fotos_clase_paths: 1 a 3 rutas locales a fotos ya
    comprimidas (ver image_utils.py)."""
    tpl = DocxTemplate(str(PLANEACION_TEMPLATE))
    ctx = dict(contexto)
    ancho = _ANCHO_FOTO_CLASE_MM.get(len(fotos_clase_paths), _IMG_WIDTH_CHICA_MM)
    ctx["fotos_clase"] = [InlineImage(tpl, p, width=Mm(ancho)) for p in fotos_clase_paths]
    tpl.render(ctx)

    ruta_salida = Path(ruta_salida)
    tpl.save(str(ruta_salida))
    _anexar_historial(ruta_salida, contexto.get("historial"))
    return ruta_salida


def generar_informe_mensual_docx(contexto: dict, ruta_salida: str | Path) -> Path:
    """contexto: el dict que devuelve api_client.generar_informe_mensual
    (incluye encuentros[].foto_base64 y firma_base64 como base64 crudo,
    que acá se convierten a InlineImage)."""
    tpl = DocxTemplate(str(INFORME_TEMPLATE))
    ctx = dict(contexto)

    ctx["encuentros"] = [
        {
            "nro": e["nro"],
            "asistencia": e.get("asistencia", ""),
            "foto": _imagen_desde_base64(tpl, e.get("foto_base64"), _IMG_WIDTH_CHICA_MM),
        }
        for e in contexto.get("encuentros", [])
    ]
    ctx["firma"] = _imagen_desde_base64(tpl, contexto.get("firma_base64"), 35)

    if ctx.get("es_directivo"):
        ctx["evidencias_gestion"] = [
            dict(ev, foto=_imagen_desde_base64(tpl, ev.get("foto_base64"), _IMG_WIDTH_CHICA_MM))
            for ev in contexto.get("evidencias_gestion", [])
        ]

    tpl.render(ctx)

    ruta_salida = Path(ruta_salida)
    tpl.save(str(ruta_salida))
    _anexar_historial(ruta_salida, contexto.get("historial"))
    return ruta_salida


def generar_certificado_pago_docx(contexto: dict, ruta_salida: str | Path) -> Path:
    """Certificado mensual de horas de docencia para pago (ver
    Certificados.js#generar_certificado_pago) — solo administradores. Sin
    historial de revisión: no es un documento que se apruebe o devuelva."""
    tpl = DocxTemplate(str(CERTIFICADO_PAGO_TEMPLATE))
    tpl.render(dict(contexto))

    ruta_salida = Path(ruta_salida)
    tpl.save(str(ruta_salida))
    return ruta_salida


def generar_informe_asistencia_docx(contexto: dict, ruta_salida: str | Path) -> Path:
    """Informe consolidado de asistencia de todos los cursos de un mes (ver
    Asistencia.js#generar_informe_asistencia) — solo administradores, para
    mandar a la Secretaría de Educación. Sin historial de revisión: es un
    reporte, no un documento que se aprueba o se devuelve.

    No usa plantilla: el documento se arma en informe_asistencia_docx (una
    página por curso, con firmas)."""
    return informe_asistencia_docx.generar(contexto, ruta_salida)


def generar_informe_asistencia_curso_docx(
    mes_nombre: str, anio: str, fecha_emision: str, curso: dict, ruta_salida: str | Path
) -> Path:
    """Igual que generar_informe_asistencia_docx pero para UN SOLO curso
    (ver Asistencia.js: cada elemento de `cursos`), para descargar/firmar
    por separado en vez del documento consolidado de todos los cursos."""
    contexto = {"mes_nombre": mes_nombre, "anio": anio, "fecha_emision": fecha_emision, "cursos": [curso]}
    return informe_asistencia_docx.generar(contexto, ruta_salida, con_anexos=False)


def generar_informe_gestion_docx(contexto: dict, ruta_salida: str | Path) -> Path:
    """Informe mensual de un directivo sin curso (ver InformesGestion.js).
    Es un documento aparte del docente: solo gestión y cuenta de cobro con
    la hora directiva."""
    tpl = DocxTemplate(str(INFORME_GESTION_TEMPLATE))
    ctx = dict(contexto)
    ctx["firma"] = _imagen_desde_base64(tpl, contexto.get("firma_base64"), 35)
    tpl.render(ctx)

    ruta_salida = Path(ruta_salida)
    tpl.save(str(ruta_salida))
    _anexar_historial(ruta_salida, contexto.get("historial"))
    return ruta_salida
