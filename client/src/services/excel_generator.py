"""Reporte de inasistencias acumuladas, en Excel (pedido de Miguel, aparte
del informe de asistencia mensual en PDF: esto es de TODAS las clases que
el curso haya dictado, no de un mes puntual, y sin necesidad de formato
para firmar — una tabla simple, en Excel, para que la revisen directo).
Ver Asistencia.js#generar_reporte_inasistencias para cómo se arma el
contexto.

Una pestaña por curso activo, cada una con la tabla de sus estudiantes:
total de clases/asistencias/faltas, con la fila resaltada si faltó más de
`umbral_alerta` veces. No hay columna de fecha de inscripción: se sacó
porque Inscripciones.creado_en no es confiable (hay clases dictadas antes
de que la inscripción quedara registrada) y mostrarla solo confundiría.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_ROJO_ALERTA = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
_GRIS_ENCABEZADO = PatternFill(start_color="FFE8E8E8", end_color="FFE8E8E8", fill_type="solid")

_COLUMNAS = ["Estudiante", "Clases dictadas", "Asistencias", "Inasistencias", "Alerta"]
_ANCHOS = [32, 15, 12, 13, 10]

# Excel prohíbe \ / ? * [ ] : en el nombre de una hoja y lo corta a 31
# caracteres -- sin esto, un curso con "/" en el nombre (pasa con horarios
# tipo "Inglés 8/9 años") rompía el archivo entero al guardarlo.
_CARACTERES_INVALIDOS = re.compile(r'[\\/?*\[\]:]')


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    limpio = _CARACTERES_INVALIDOS.sub("-", nombre).strip() or "Curso"
    base = limpio[:31]
    candidato = base
    sufijo = 2
    # Dos cursos truncados al mismo nombre de 31 caracteres, o dos cursos
    # con el mismo nombre real: se numeran en vez de que uno pise al otro.
    while candidato.lower() in usados:
        recorte = 31 - len(f" ({sufijo})")
        candidato = f"{base[:recorte]} ({sufijo})"
        sufijo += 1
    usados.add(candidato.lower())
    return candidato


def _encabezado(ws: Worksheet, fila: int) -> None:
    for col, (titulo, ancho) in enumerate(zip(_COLUMNAS, _ANCHOS), start=1):
        celda = ws.cell(row=fila, column=col, value=titulo)
        celda.font = Font(bold=True)
        celda.fill = _GRIS_ENCABEZADO
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _hoja_curso(ws: Worksheet, curso: dict, fecha_emision: str, umbral: int) -> None:
    ws.cell(row=1, column=1, value=curso.get("nombre", "")).font = Font(bold=True, size=13)
    ws.cell(
        row=2, column=1,
        value=f"Docente: {curso.get('docente', '')}    Núcleo: {curso.get('nucleo', '')}",
    ).font = Font(italic=True, color="FF666666")
    ws.cell(
        row=3, column=1,
        value=f"Fecha de emisión: {fecha_emision}    ·    Resaltados: más de {umbral} inasistencias",
    ).font = Font(italic=True, color="FF999999")

    fila = 5
    _encabezado(ws, fila)
    fila += 1

    estudiantes = curso.get("estudiantes", [])
    if not estudiantes:
        ws.cell(row=fila, column=1, value="(sin estudiantes inscritos)").font = Font(italic=True, color="FF999999")
        fila += 1
    for est in estudiantes:
        valores = [
            est.get("nombre", ""),
            est.get("total_clases", 0), est.get("total_asistio", 0), est.get("total_falto", 0),
            "Sí" if est.get("alerta") else "",
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            if est.get("alerta"):
                celda.fill = _ROJO_ALERTA
                if col == len(valores):
                    celda.font = Font(bold=True, color="FFCC0000")
        fila += 1

    ws.freeze_panes = "A6"


def generar_reporte_inasistencias_xlsx(contexto: dict, ruta_salida: str | Path) -> Path:
    cursos = contexto.get("cursos", [])
    fecha_emision = contexto.get("fecha_emision", "")
    umbral = contexto.get("umbral_alerta", 3)

    wb = Workbook()
    usados: set[str] = set()

    if not cursos:
        ws = wb.active
        ws.title = "Inasistencias"
        ws.cell(row=1, column=1, value="No hay cursos activos con estudiantes inscritos.")
    else:
        primero = True
        for curso in cursos:
            ws = wb.active if primero else wb.create_sheet()
            ws.title = _nombre_hoja(curso.get("nombre", "Curso"), usados)
            _hoja_curso(ws, curso, fecha_emision, umbral)
            primero = False

    ruta_salida = Path(ruta_salida)
    wb.save(str(ruta_salida))
    return ruta_salida
